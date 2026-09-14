import json
import re
import sys
from pptx import Presentation

prs = Presentation("H&E空间转录组生成研究_阶段汇报_庹力元.pptx")


def col(fmt):
    try:
        if fmt.type is not None:
            s = str(fmt.type)
            if s.endswith("SOLID (1)"):
                return "rgb:" + str(fmt.fore_color.rgb)
            if s.endswith("GRADIENT (2)"):
                return "gradient:" + ",".join(str(x.color.rgb) for x in fmt.gradient_stops)
    except Exception as exc:
        return "err:" + str(exc)
    return None


slide = prs.slides[13]
for shape in slide.shapes:
    if shape.shape_type == 13:
        continue
    info = {"id": shape.shape_id, "fill": col(shape.fill)}
    if shape.has_text_frame:
        runs = []
        for p in shape.text_frame.paragraphs:
            for r in p.runs:
                c = None
                try:
                    if r.font.color.type is not None:
                        c = str(r.font.color.rgb)
                except Exception:
                    pass
                runs.append({"t": r.text[:14], "font": r.font.name,
                             "size": r.font.size.pt if r.font.size else None,
                             "b": r.font.bold, "color": c})
        info["runs"] = runs
    xml = shape._element.xml
    m = re.search(r'prstGeom prst="(\w+)"', xml)
    adj = re.search(r'<a:gd name="adj" fmla="val ([0-9]+)"', xml)
    if m:
        info["geom"] = m.group(1)
    if adj:
        info["adjust"] = adj.group(1)
    print(json.dumps(info, ensure_ascii=False))
