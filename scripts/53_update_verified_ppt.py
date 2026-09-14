"""Update the existing 11-slide deck in place, retaining its first five slides.

Slides 7-8 explain verified causes and ResSAT reproduction. Other result slides
are updated to the completed fixed-HEG benchmark. Requires a complete benchmark.
"""
from pathlib import Path
from datetime import datetime
import importlib.util
import json
import shutil
import sys
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'results/verified_20260913'
PPT = ROOT / 'H&E空间转录组生成研究_统一Benchmark汇报_庹力元.pptx'
FIG = ROOT / 'figures/verified_20260913'
FIG.mkdir(exist_ok=True, parents=True)
spec = importlib.util.spec_from_file_location('ppt_style', ROOT / 'scripts/48_update_unified_ppt.py')
style = importlib.util.module_from_spec(spec)
spec.loader.exec_module(style)
txt, rect, pic = style.txt, style.rect, style.picture
RED, GOLD, INK, GRAY = style.RED, style.GOLD, style.INK, style.GRAY
plt.rcParams.update({'font.family': 'Microsoft YaHei', 'axes.unicode_minus': False})


def clear_slide(slide):
    for shape in list(slide.shapes):
        el = shape._element
        el.getparent().remove(el)


def note_frame(slide):
    notes = slide.notes_slide
    if notes.notes_text_frame is None:
        for placeholder in notes.part.notes_master.placeholders:
            if placeholder.placeholder_format.type == PP_PLACEHOLDER.BODY:
                notes.shapes.clone_placeholder(placeholder)
                break
    if notes.notes_text_frame is None:
        notes.shapes._spTree.add_placeholder(notes.shapes._next_shape_id,
            'Notes Placeholder', PP_PLACEHOLDER.BODY, 'horz', 'full', 1)
    assert notes.notes_text_frame is not None
    return notes.notes_text_frame


def header(slide, page, title, subtitle, source, notes):
    clear_slide(slide)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string('FFFFFF')
    txt(slide, title, .48, .26, 12.0, .55, 26, RED, True)
    rect(slide, .5, .95, 12.3, .025, GOLD)
    txt(slide, subtitle, .52, 1.07, 12.22, .63, 15, GRAY)
    txt(slide, f'{page:02d} / 11', 12.0, .28, .8, .28, 10, GRAY, align=PP_ALIGN.RIGHT)
    txt(slide, source, .53, 7.08, 12.1, .25, 9, GRAY)
    note_frame(slide).text = notes


def card(slide, x, y, w, h, title, body, fill='F7F4EF', size=16):
    rect(slide, x, y, w, h, fill, rounded=True)
    txt(slide, title, x+.17, y+.12, w-.34, .45, 18, RED, True)
    txt(slide, body, x+.17, y+.65, w-.34, h-.7, size, INK)


def make_maps(data, methods, gene):
    j = list(data['genes']).index(gene)
    xy = data['coordinates_xy']
    names = ['实测真值'] + list(methods)
    matrices = [data['truth']] + [data[n] for n in methods]
    vmin, vmax = np.quantile(data['truth'][:, j], [.01, .99])
    fig, axes = plt.subplots(2, 4, figsize=(13, 4.7), layout='constrained')
    truth = data['truth'][:, j]
    stats = {}
    for ax, name, matrix in zip(axes.flat, names, matrices):
        values = matrix[:, j]
        sc = ax.scatter(xy[:, 0], xy[:, 1], c=values, s=3.3, cmap='viridis', vmin=vmin, vmax=vmax, rasterized=True)
        pcc = float(np.corrcoef(truth, values)[0, 1])
        stats[name] = pcc
        display = 'ST-Net' if name.startswith('ST-Net') else name
        ax.set_title(display if name == '实测真值' else f'{display}  PCC={pcc:.3f}', fontsize=12)
        ax.set_aspect('equal'); ax.invert_yaxis(); ax.axis('off')
    fig.colorbar(sc, ax=axes, shrink=.65, pad=.015, label='面板相对表达 log1p')
    fig.savefig(FIG / f'{gene}_fixed_HEG.png', dpi=180, facecolor='white')
    plt.close(fig)
    return stats


def main():
    results = json.loads((OUT / 'benchmark.json').read_text(encoding='utf8'))
    assert not results['pending'] and len(results['methods']) == 7
    panel = json.loads((OUT / 'panel_audit.json').read_text(encoding='utf8'))
    paper = json.loads((OUT / 'ressat_recomputed.json').read_text())
    fresh = json.loads((OUT / 'fresh_checkpoint_verification.json').read_text())
    assert fresh['SA']['match'] and fresh['SP']['match']
    methods = results['methods']
    data = dict(np.load(OUT / 'benchmark_predictions.npz'))
    # Highest and second-highest train expression: examples chosen without PCC.
    examples = panel['primary_top50'][:2]
    map_stats = {g: make_maps(data, methods, g) for g in examples}
    backup = OUT / 'original_ppt_backup'
    backup.mkdir(exist_ok=True)
    backup_file = backup / PPT.name
    if not backup_file.exists(): shutil.copy2(PPT, backup_file)
    prs = Presentation(PPT)
    assert len(prs.slides) == 11
    for slide in list(prs.slides)[:5]:
        for shape in slide.shapes:
            if not shape.has_text_frame: continue
            for p in shape.text_frame.paragraphs:
                whole_paragraph = {
                    '回归：Ridge / MLP': '回归：Ridge/MLP/ST-Net',
                    '回归：Ridge / MLP / ST-Net': '回归：Ridge/MLP/ST-Net',
                    '本地：SA/SP 四张原图、计数、': '本地：SA/SP 数据与模型均可用',
                    'arrays/C73_D1.npz': '../gse240429_heg/arrays/C73_D1.npz',
                }
                if p.text in whole_paragraph and p.runs:
                    value = whole_paragraph[p.text]
                    p.runs[0].text = value
                    for remaining in p.runs[1:]: remaining.text = ''
                if p.text == '回归：Ridge/MLP/ST-Net':
                    for run in p.runs: run.font.size = Pt(16)
                for run in p.runs:
                    run.text = (run.text.replace('六种方法', '七种方法').replace('六法', '七法')
                                .replace('200 个随机基因', '200 个高表达基因')
                                .replace('固定 200 基因', '训练集选定 200 个高表达基因')
                                .replace('200 个基因只从训练切片筛选', '200 个高表达基因只从训练切片筛选')
                                .replace('本地：SA/SP 四张原图、计数、', '本地：SA/SP 数据与模型均可用')
                                .replace('右侧目录前缀：data/processed/gse240429/', '图像目录保持不变；新标签来自 gse240429_heg/'))
        # Keep the original visuals; add current benchmark provenance to notes.
        notice = '2026-09-13：主基准改用 data/processed/gse240429_heg；仅训练 A1+B1 选 HEG200/HEG50。图像缓存仍为 gse240429/image。新增真实 DenseNet121 ST-Net 适配，MLP 单独列示。'
        if notice not in note_frame(slide).text:
            note_frame(slide).text += '\n' + notice

    s = prs.slides[5]
    header(s, 6, '统一高表达基因 Benchmark：已完成七法预测',
           'A1+B1 训练，C1 选择配置，D1 测试；同一组 HEG200 / HEG50，真值不做 PCA 重建。',
           '来源：results/verified_20260913/benchmark.csv；基因选择、行列顺序和预测文件 SHA256 均已核验。',
           json.dumps(results['protocol'], ensure_ascii=False) + '\nHEG50 按训练集全转录组归一化平均表达排序，不能按测试 PCC 排序。七法共同目标为200基因面板内相对表达，MAE也在此单位；不是原始计数误差。')
    cols = [.65, 4.5, 6.6, 8.6, 10.25]
    for x, w, label in zip(cols, [3.7, 1.9, 1.9, 1.5, 2.0], ['方法', 'HEG200 PCC ↑', 'HEG50 PCC ↑', 'MAE ↓', '实现']):
        txt(s, label, x, 1.91, w, .38, 15, RED, True)
    rows = sorted(methods.items(), key=lambda kv: kv[1]['HEG50_PCC'], reverse=True)
    impl = {'Image Ridge':'冻结 ResNet18', 'MLP':'冻结 ResNet18', 'BLEEP':'对比 / 检索适配',
            'ResSAT':'官方主体适配', 'GenAR':'自回归 + 期望解码', 'Stem':'官方 DiT 单卡适配', 'ST-Net (DenseNet121)':'DenseNet121 微调'}
    for i, (name, d) in enumerate(rows):
        y = 2.43 + i*.43
        rect(s, .6, y-.03, 12.08, .42, 'F7F4EF' if i % 2 == 0 else 'FFFFFF')
        for x, w, val in zip(cols, [3.7, 1.9, 1.9, 1.5, 2.0], [name, f'{d["HEG200_PCC"]:.4f}', f'{d["HEG50_PCC"]:.4f}', f'{d["MAE"]:.3f}', impl[name]]):
            txt(s, val, x, y, w, .35, 14, INK, i == 0)
    card(s, .6, 5.63, 12.08, 1.12, '这张表能说明什么？',
         '同一高表达基因任务上的适配结果。编码器与训练预算各异，且仅 1 位供者、1 张测试切片，不能据此宣称普适排名。', 'EDF3F8', 14)

    s = prs.slides[6]
    header(s, 7, '为什么 PCC 只有 0.0x？已证实的问题与修正',
           '先纠正基因选择和计分方式，再讨论模型性能；原 PPT 中“不是模型写错 / 主因就是域偏移”的断言已撤回。',
           '证据：panel_audit.json、tiff_patch_audit.json、ressat_recomputed.json、各方法 selection.json（均在 results/）。',
           '旧面板按训练表达方差排序，不是随机基因，也不是训练集全转录组表达最高的200基因。高变不自动等于高表达。零值下降不单独证明全部PCC差距的因果。当前人肝图像缓存是uint8、抽样128块与原图一致，历史官方float图像问题已修复。生成式修正的贡献必须看同基因同单位消融。域偏移是待进一步验证的解释。')
    card(s, .6, 1.92, 5.92, 2.16, '① 基因面板不符合高表达要求',
         '旧 200 基因与训练集 HEG200 重叠：0 个。\nD1 零值：24.70% → 3.45%。\n修正：仅 A1+B1 排序，固定 HEG200 / HEG50。', 'F9EDEC', 16)
    card(s, 6.72, 1.92, 5.98, 2.16, '② 不同的真值与归一化尺度被混比',
         '原计数、全转录组归一化、面板归一化、PCA 重建是不同目标。\n修正：主表统一单位；PCA 结果只作补充。', 'EDF3F8', 16)
    card(s, .6, 4.29, 5.92, 2.27, '③ 生成式实验未完成，且有实现问题',
         'GenAR 上次训练中断，未完成 HEG 推理。\nStem 旧采样截到 [-1,1]；旧 EMA 仍含 64% 初值。\n修正：补齐推理 / 重训，C1 选择解码和采样。', 'EDF5EF', 15)
    card(s, 6.72, 4.29, 5.98, 2.27, '④ 比较与汇报也需要纠错',
         'MLP 不能标成完整 ST-Net；本轮补 DenseNet121。\n旧“0.087→0.236”同时换基因和计分方式。\n修正：不再当作同任务的模型提升。', 'F7F4EF', 15)

    s = prs.slides[7]
    header(s, 8, 'ResSAT 关键结果核验：高表达基因达到论文量级',
           '本轮从既有 checkpoint 重新推理并核验；SA 使用作者示例，SP 从原始 10x 数据按 Methods 重建。',
           '论文：Genome Biology 2026，Tables 1–2；本地：fresh_checkpoint_verification.json / ressat_recomputed.json。',
           'https://link.springer.com/content/pdf/10.1186/s13059-026-04168-x_reference.pdf\n论文平均来自5次重复；本地为seed42单次，不等于完整多seed复现。SA预处理使用作者文件，SP是Methods重建，不能称逐字节复现。论文2000HVG/50HEG成绩都在PCA/Harmony重建空间，HEG是该HVG面板内按观察表达选的50个。联合预处理使用各section表达，是转导式口径，不能代表完全未见切片/患者。')
    txt(s, '指标 / 数据集', .75, 1.94, 5.0, .4, 17, RED, True)
    txt(s, '论文', 6.55, 1.94, 2.0, .4, 17, RED, True)
    txt(s, '本地重算', 9.5, 1.94, 2.5, .4, 17, RED, True)
    vals = [('2,000 HVG · SA', paper['SA']['paper_HVG2000'], paper['SA']['reproduced_HVG2000']),
            ('2,000 HVG · SP', paper['SP']['paper_HVG2000'], paper['SP']['reproduced_HVG2000']),
            ('50 HEG · SA', paper['SA']['paper_HEG50'], paper['SA']['reproduced_HEG50']),
            ('50 HEG · SP', paper['SP']['paper_HEG50'], paper['SP']['reproduced_HEG50'])]
    for i, (label, ref, val) in enumerate(vals):
        y = 2.54 + i*.52
        rect(s, .62, y-.04, 12.05, .48, 'F7F4EF' if i < 2 else 'EDF5EF')
        for x, w, value in [(.75,5,label),(6.55,2,f'{ref:.4f}'),(9.5,2.5,f'{val:.4f}')]:
            txt(s, value, x, y, w, .39, 18, INK, i >= 2)
    card(s, .62, 4.88, 5.92, 1.84, '为什么不能拿 0.6980 要求人肝实验？',
         f'同一 SP 预测，改用未重建真值：\n2,000 HVG PCC = {paper["SP"]["strict_truth_HVG2000_defined_macro"]:.4f}；\n相同 50 HEG PCC = {paper["SP"]["strict_truth_same_HEG50"]:.4f}。', 'EDF3F8', 15)
    card(s, 6.73, 4.88, 5.92, 1.84, '复现结论与边界',
         '高表达基因已接近论文量级。\n本地是单 seed，SP 仍有预处理差异。\n人肝严格真值的成绩仍需继续提高。', 'F9EDEC', 15)

    for index, gene in zip([8, 9], examples):
        s = prs.slides[index]
        rank = examples.index(gene) + 1
        header(s, index+1, f'固定高表达基因实例：{gene}（训练表达第 {rank}）',
               '按训练表达排名预先选例，不按测试 PCC 挑最好看的基因；同一 D1 真值、同一坐标、各方法共用色标。',
               '来源：results/verified_20260913/benchmark_predictions.npz；显示范围由真值 1%–99% 分位数确定。',
               json.dumps(map_stats[gene], ensure_ascii=False) + '\n色标裁剪仅用于显示，不改变计算PCC的数值；地图平滑或颜色丰富不保证空间表达预测正确。')
        pic(s, FIG / f'{gene}_fixed_HEG.png', .51, 1.86, 12.28, 4.58)
        txt(s, '判断依据：高表达区域是否落在正确位置；整体结论以第 6 页固定 HEG200 / HEG50 的平均 PCC 为准。',
            .66, 6.56, 12.0, .4, 14, GRAY)

    s = prs.slides[10]
    best = max(methods, key=lambda n: methods[n]['HEG50_PCC'])
    header(s, 11, '当前结论：已完成的结果与仍存在的差距',
           '数据已核验，七法高表达基因对比已落盘；论文复现与严格预测任务分别报告。',
           '完整诊断：docs/05_PCC_DIAGNOSIS.md；指标、逐基因结果、预测矩阵：results/verified_20260913/。',
           '不声称所有方法已复现论文优势，不把单供者测试当作独立患者泛化。后续按患者留出、匹配病理编码器与预算，保留HEG200/HEG50训练选基因约束。')
    card(s, .62, 1.94, 12.05, 1.48, '已完成：高表达基因对比与 ResSAT 核验',
         f'本轮 D1 固定 HEG50 最好：{best}，PCC={methods[best]["HEG50_PCC"]:.4f}。\nResSAT 鼠脑论文口径 HEG50：SA 0.8803 / SP 0.8930；各方法完整预测均可复核。', 'EDF5EF', 16)
    card(s, .62, 3.64, 5.92, 2.12, '方法优势与限制',
         'Ridge / MLP：成本低，作必要基线。\nBLEEP：检索训练表达，受训练覆盖限制。\nResSAT / ST-Net：学习形态特征，需微调。\nGenAR / Stem：可生成表达，但本地适配受限。', 'F7F4EF', 15)
    card(s, 6.73, 3.64, 5.92, 2.12, '仍需解决',
         '人肝严格真值尚未达到鼠脑论文的 PCC。\n统一编码器 / 预算，增加独立患者验证。\n高表达基因选择已固定，后续不按测试成绩\n换基因、换切片或用去噪成绩替代主指标。', 'EDF3F8', 15)
    txt(s, '真正的改进应在相同基因、相同测试集、相同真值和计分方式下成立。', .72, 6.12, 11.95, .5, 21, RED, True)
    prs.core_properties.title = 'H&E 空间转录组生成研究：高表达基因 Benchmark 与 PCC 诊断'
    prs.core_properties.subject = '2026-09-13 verified results; 11 slides'
    style.renumber(prs)
    prs.save(PPT)
    # Native PPTX structure / page-boundary checks.
    check = Presentation(PPT)
    assert len(check.slides) == 11
    problems = []
    for i, slide in enumerate(check.slides, 1):
        for shape in slide.shapes:
            if i >= 6 and (shape.left < 0 or shape.top < 0 or shape.left + shape.width > check.slide_width + 1000 or shape.top + shape.height > check.slide_height + 1000):
                problems.append([i, shape.name])
    assert not problems, problems
    (OUT / 'ppt_verification.json').write_text(json.dumps(dict(file=PPT.name, slides=11,
        backup=str(backup_file.relative_to(ROOT)), examples_selected_by_training_expression=examples,
        no_pending_results=True, edited_slides=[6,7,8,9,10,11], shape_bounds_passed=True), ensure_ascii=False, indent=2), encoding='utf8')
    print(PPT, '11 slides updated in place')


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf8')
    main()
