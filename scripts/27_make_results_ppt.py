"""生成九页统一 Benchmark 汇报，复用已有预测，不训练、不下载。

运行：py -3.12 scripts/27_make_results_ppt.py
保留旧 PPT，沿用用户的白底、深红标题、金色分隔线。
评测修正：所有方法都按同一个 200 基因面板归一化；空间图共享色标。
"""
from pathlib import Path
import gzip
import hashlib
import json
import sys
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "H&E空间转录组生成研究_统一Benchmark汇报_庹力元.pptx"
ASSETS = ROOT / "figures/benchmark_ppt"
RED, GOLD, INK, GRAY = "B71C1C", "B08D57", "262626", "666666"
LIGHT, BLUE, GREEN, PINK = "F7F4EF", "EDF3F8", "EDF5EF", "F9EDEC"
METHODS = [
    ("Image Ridge", "unified_image_baselines/image_ridge_C73_D1.npz", "predicted"),
    ("MLP 基线", "unified_image_baselines/stnet_style_mlp_C73_D1.npz", "predicted"),
    ("BLEEP", "unified_image_baselines/bleep_C73_D1.npz", "predicted"),
    ("ResSAT", "ressat_unified/unified_batch16_predictions.npz", "predicted"),
    ("GenAR", "genar_gse240429/C73_D1_predictions.npz", "predicted_counts"),
    ("Stem", "stem_adapted/stem_samples_D1.npz", "mean_prediction"),
]
SOURCES = {
    "BLEEP": "https://proceedings.neurips.cc/paper_files/paper/2023/hash/df656d6ed77b565e8dcdfbf568aead0a-Abstract-Conference.html",
    "GEO": "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE240429",
    "ResSAT": "https://link.springer.com/article/10.1186/s13059-026-04168-x",
    "ST-Net": "https://www.nature.com/articles/s41551-020-0578-x",
    "Stem": "https://openreview.net/forum?id=FtjLUHyZAO",
    "GenAR": "https://arxiv.org/abs/2510.04315",
    "HEST": "https://github.com/mahmoodlab/HEST",
}
plt.rcParams.update({"font.family": "Microsoft YaHei", "axes.unicode_minus": False})


def panel_normalize(values):
    """非负化 → 同一 200 基因合计缩放至 10000 → log1p。

    输入是计数或相对丰度；后者的整行比例因子在本步骤消去。
    预测只用预测自己的行和，不借用测试标签的总 UMI。面板全零行原样保留。
    """
    values = np.maximum(np.asarray(values, dtype=np.float64), 0)
    return np.log1p(10000 * values / np.maximum(values.sum(1, keepdims=True), 1e-12))


def gene_pcc(truth, prediction):
    """每基因跨所有位置计算 PCC；常数预测的 PCC 无定义，不伪装成正相关。"""
    x, y = truth - truth.mean(0), prediction - prediction.mean(0)
    den = np.sqrt((x*x).sum(0) * (y*y).sum(0))
    valid = (np.ptp(truth, axis=0) > 1e-10) & (np.ptp(prediction, axis=0) > 1e-10)
    return np.divide((x*y).sum(0), den, out=np.full(truth.shape[1], np.nan), where=valid)


def load_and_audit():
    """检查行列对应与真实标签，统一单位后重算，不修改预测文件。"""
    with np.load(ROOT / "data/processed/gse240429/arrays/C73_D1.npz") as data:
        raw, original = data["raw_counts"], data["log_normalized"]
        genes, xy = data["genes"], data["coordinates_xy"]
    truth = panel_normalize(raw)
    predictions, stats, checks = {}, {}, []
    for name, relative, key in METHODS:
        path = ROOT / "results" / relative
        with np.load(path) as data:
            is_count = name == "GenAR"
            assert np.array_equal(data["gene_names" if is_count else "genes"], genes)
            # GenAR 没有坐标字段；逐元素计数和基因顺序相同，核对其标签行序。
            assert np.array_equal(data["true_counts" if is_count else "true"], raw if is_count else original)
            if not is_count:
                assert np.array_equal(data["coordinates_xy"], xy)
            values = data[key].astype(np.float64)
        assert values.shape == truth.shape == (2265, 200) and np.isfinite(values).all()
        prediction = panel_normalize(values if is_count else np.expm1(values))
        correlations = gene_pcc(truth, prediction)
        assert np.isfinite(correlations).all(), f"{name}: 存在无定义 PCC，需另报"
        predictions[name] = prediction
        stats[name] = {
            "pcc": float(correlations.mean()), "mae": float(np.abs(truth-prediction).mean()),
            "rmse": float(np.sqrt(np.square(truth-prediction).mean())),
            "negative_fraction_before_clipping": float(np.mean(values < 0)),
            "per_gene_pcc": dict(zip(genes.tolist(), correlations.tolist())),
        }
        checks.append({"model": name, "file": relative, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                       "gene_order": "pass", "truth_and_row_order": "pass", "shape_and_finite": "pass"})
    with np.load(ROOT / "results/nonimage_baselines/train_gene_mean_C73_D1.npz") as data:
        assert np.array_equal(data["true"], original)
        mean_prediction = panel_normalize(np.expm1(data["predicted"].astype(np.float64)))
    assert np.ptp(mean_prediction, axis=0).max() < 1e-10
    audit = {
        "protocol": "GSE240429: A1+B1 train; C1 validation; D1 test; fixed 200 genes",
        "evaluation_space": "log1p(10000 * nonnegative abundance / sum of SAME 200 genes within each spot)",
        "normalization_change": "Old table mixed GenAR panel-total normalization with all-gene-total normalization for other models. All models now use panel-total normalization. No retraining.",
        "prediction_conversion": "GenAR: raw predicted_counts. Other models: max(expm1(saved log prediction),0). Truth: raw 200-gene counts. Normalize each matrix independently.",
        "no_test_labels_used_for_prediction_conversion": True,
        "truth_zero_panel_rows_retained": int(np.sum(raw.sum(1) == 0)),
        "stats": stats, "checks": checks,
        "mean_baseline": {"pcc": None, "mae": float(np.abs(truth-mean_prediction).mean())},
        "limitations": ["One sample, slice-heldout, not patient-heldout.",
                        "Local adaptations use different backbones and training budgets.",
                        "Panel-relative expression is not absolute counts or full-transcriptome recovery.",
                        "GLUL and IGKC are illustrative; report all 200 genes, not test-selected top-k."],
    }
    (ROOT / "results/benchmark_ppt_audit.json").write_text(json.dumps(audit, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return truth, genes, xy, predictions, audit


def make_figures(truth, genes, xy, predictions, audit):
    """从真实数据重画，不裁旧比较图，不平滑，不独立拉伸各模型色标。"""
    ASSETS.mkdir(parents=True, exist_ok=True)
    for section in ["A1", "B1", "C1", "D1"]:
        path = next((ROOT/"data/raw/GSE240429").glob(f"*C73{section}_tissue_lowres_image.png.gz"))
        with gzip.open(path, "rb") as stream:
            Image.open(stream).convert("RGB").save(ASSETS/f"he_{section}.png")
    patches = np.load(ROOT/"data/processed/gse240429/image/C73_D1_patches_uint8.npy", mmap_mode="r")
    assert len(patches) == len(truth)
    Image.fromarray(patches[100]).save(ASSETS/"example_patch.png")
    for gene in ["GLUL", "IGKC"]:
        index = list(genes).index(gene)
        entries = [("实测真值", truth)] + list(predictions.items())
        maximum = max(float(values[:, index].max()) for _, values in entries)
        fig, axes = plt.subplots(2, 4, figsize=(13.6, 5.55))
        fig.subplots_adjust(left=.015, right=.985, bottom=.015, top=.91, wspace=.08, hspace=.36)
        for ax, (name, values) in zip(axes.flat, entries):
            dots = ax.scatter(xy[:, 0], xy[:, 1], c=values[:, index], s=4.3, cmap="magma",
                              vmin=0, vmax=maximum, linewidths=0, rasterized=True)
            ax.set(xlim=(xy[:,0].min()-400, xy[:,0].max()+400),
                   ylim=(xy[:,1].max()+400, xy[:,1].min()-400), aspect="equal")
            ax.axis("off")
            subtitle = "参考答案" if name == "实测真值" else f"PCC = {audit['stats'][name]['per_gene_pcc'][gene]:.3f}"
            ax.set_title(name+"\n"+subtitle, fontsize=15, pad=3, color="#262626")
        axes[1,3].axis("off")
        colorbar = fig.colorbar(dots, cax=fig.add_axes([.814, .17, .024, .20]))
        colorbar.ax.tick_params(labelsize=13)
        colorbar.set_label("面板内归一化表达", fontsize=13)
        fig.text(.793, .435, "统一色标", fontsize=16, weight="bold")
        fig.text(.793, .052, "亮：表达高\n暗：表达低", fontsize=14)
        fig.savefig(ASSETS/f"predictions_{gene}.png", dpi=190, facecolor="white")
        plt.close(fig)
    idx = list(genes).index("GLUL")
    fig, ax = plt.subplots(figsize=(3,3), layout="constrained")
    ax.scatter(xy[:,0], xy[:,1], c=truth[:,idx], s=7, linewidths=0, cmap="magma")
    ax.invert_yaxis(); ax.set_aspect("equal"); ax.axis("off")
    fig.savefig(ASSETS/"truth_GLUL.png", dpi=180); plt.close(fig)


def rect(slide, x, y, w, h, fill=LIGHT, rounded=False):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                    Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill); shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = .06
    return shape


def txt(slide, value, x, y, w, h, size=18, ink=INK, bold=False, align=PP_ALIGN.LEFT, url=None):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(.015)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for i, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line; p.font.name = "Microsoft YaHei"; p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(ink); p.font.bold = bold; p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(4); p.line_spacing = 1.12
        if url:
            p.runs[0].hyperlink.address = url
    return shape


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w/iw, h/ih)
    return slide.shapes.add_picture(str(path), Inches(x+(w-iw*scale)/2), Inches(y+(h-ih*scale)/2),
                                    width=Inches(iw*scale), height=Inches(ih*scale))


def new_slide(prs, heading, subtitle, source, notes=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    txt(s, heading, .48, .26, 12.25, .51, 27, RED, True)
    rect(s, .5, .94, 12.3, .025, GOLD)
    txt(s, subtitle, .51, 1.07, 12.24, .57, 16, GRAY)
    txt(s, source, .5, 7.12, 11.86, .22, 9, GRAY)
    txt(s, f"{len(prs.slides):02d} / 09", 12.0, 7.11, .79, .24, 10, GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes+"\n\n资料来源：\n"+"\n".join(f"{k}: {v}" for k,v in SOURCES.items())
    return s


def card(slide, heading, body, x, y, w, h, fill=LIGHT, size=17):
    rect(slide, x, y, w, h, fill, True)
    txt(slide, heading, x+.19, y+.14, w-.38, .46, 20, RED, True)
    txt(slide, body, x+.19, y+.72, w-.38, h-.83, size)


def make_deck(truth, genes, xy, predictions, audit):
    prs = Presentation(); prs.slide_width = Inches(13.333333); prs.slide_height = Inches(7.5)
    stats = audit["stats"]
    s = new_slide(prs, "从 H&E 图像预测空间基因表达", "统一 Benchmark 实验汇报  ·  庹力元  ·  2026.09", "真实图像与表达：GSE240429 / C73_D1；图中的表达示例是实测值，不是模型预测。",
                  "任务输出是位置×基因矩阵，不是染色图片。图中patch为D1第101个位置的真实缓存。spot是测量区域，通常包含多个细胞，不等于单细胞。GLUL图用真实表达演示输出的形式。")
    for x,label in [(.65,"① 输入：H&E + 位置"),(4.9,"② 模型预测"),(8.7,"③ 输出：基因表达")]:
        txt(s,label,x,1.8,3.95,.42,22,RED,True)
    picture(s,ASSETS/"he_D1.png",.6,2.35,2.8,2.75)
    picture(s,ASSETS/"example_patch.png",3.1,3.2,1.3,1.3)
    txt(s,"224 × 224 图像块",2.96,4.6,1.72,.42,13)
    txt(s,"→",4.42,3.24,.5,.62,30,RED,True)
    card(s,"同一张图，六种方法","回归：Ridge / MLP\n检索：BLEEP\n空间建模：ResSAT\n生成：GenAR / Stem",4.98,2.47,3.3,2.72,BLUE,17)
    txt(s,"→",8.28,3.24,.5,.62,30,RED,True)
    picture(s,ASSETS/"truth_GLUL.png",8.96,2.28,3.12,2.88)
    txt(s,"GLUL 实测示例：一列 → 一张热图",8.55,5.0,4.1,.46,15)
    rect(s,.66,5.61,12.0,1.12,LIGHT,True)
    txt(s,"最终预测：2,265 个位置 × 200 个基因 = 453,000 个表达值",.92,5.79,11.5,.42,23,RED,True)
    txt(s,"训练时有实测表达作标签；测试时只提供图像与位置，实测表达留作评分。",.92,6.3,11.5,.33,17)

    s = new_slide(prs,"统一 Benchmark 用的是什么数据？","GSE240429：同一人肝样本 C73 的 4 张连续切片；10x Visium，每个测量区域（spot）直径约 55 μm。","来源：BLEEP §3.1；data/processed/gse240429/manifest.json。",
                  "本项目定义的切片留出测试，不是官方排行榜，也不是跨患者测试。四切片共9269 spots。A1+B1训练，C1选模型，D1最终评价。200基因仅用训练集选择。评测时统一换算面板内相对表达；未改动旧训练目标和权重。")
    for i,(section,count,role,fill) in enumerate([("A1",2378,"训练",BLUE),("B1",2349,"训练",BLUE),("C1",2277,"验证",LIGHT),("D1",2265,"测试",PINK)]):
        x=.64+i*3.12
        rect(s,x,1.91,2.94,3.14,fill,True)
        txt(s,f"{section} · {role}",x+.12,2.05,2.7,.4,22,RED,True,PP_ALIGN.CENTER)
        picture(s,ASSETS/f"he_{section}.png",x+.38,2.51,2.18,1.91)
        txt(s,f"{count:,} 个位置",x+.15,4.56,2.64,.38,18,align=PP_ALIGN.CENTER)
    txt(s,"训练 4,727  +  验证 2,277  +  测试 2,265  =  共 9,269 个位置",.68,5.29,12.0,.43,22,RED,True)
    txt(s,"一个样本 = 图像块 + 坐标 + 200 个实测基因值，用同一位置编号（barcode）对齐。\n200 个基因只从训练切片筛选；整张切片划分，测试位置不参与训练。",.68,5.94,11.99,.86,18)

    s = new_slide(prs,"下载的 H&E 原图在哪里？数据怎样配成一对？","项目根目录：D:\\Code\\H&E Generation\\    ｜    下列路径均相对于这个目录。","文件已现场核对；TIFF 名称带超链接。展示用缩略图来自相应切片的 Space Ranger 输出。",
                  "原始压缩图：data/raw/GSE240429/GSM7697868–71_GEX_C73_[A1–D1]_Merged.tiff.gz。坐标CSV、barcode、features和matrix.mtx位于同一raw目录。右侧image/和arrays/的共同前缀是data/processed/gse240429/。arrays/C73_D1.npz含raw_counts、log_normalized、coordinates_xy、barcodes、genes。patch缓存与数组保持同一行序。")
    rect(s,.64,1.91,6.9,4.86,LIGHT,True)
    txt(s,"已解压的四张染色原图",.86,2.07,6.4,.43,22,RED,True)
    txt(s,"data/processed/gse240429/tiff/",.86,2.67,6.37,.38,20,bold=True)
    for i,section in enumerate(["A1","B1","C1","D1"]):
        path=ROOT/f"data/processed/gse240429/tiff/C73_{section}.tif"
        assert path.is_file()
        txt(s,f"C73_{section}.tif     {path.stat().st_size/1e9:.2f} GB",.93,3.25+i*.48,5.86,.4,19,url=path.as_uri())
    txt(s,"下载压缩包：data/raw/GSE240429/\n原图较大；PPT 使用同切片的低分辨率预览。",.87,5.56,6.35,.86,16,GRAY)
    txt(s,"右侧目录前缀：data/processed/gse240429/",7.83,1.55,4.88,.30,12,GRAY)
    card(s,"① H&E → 图像块","按坐标从 TIFF 裁 224 × 224 像素\nimage/C73_D1_patches_uint8.npy",7.82,1.91,4.86,1.42,BLUE,14.5)
    card(s,"② 实测计数 + 坐标 → 标签","matrix.mtx + barcode + 位置 CSV\n按位置编号一一匹配，不能随意排序。",7.82,3.53,4.86,1.48,GREEN,15)
    card(s,"③ 整理后 → 可直接训练的数据","arrays/C73_D1.npz\n表达 [2265, 200]；坐标 [2265, 2]",7.82,5.15,4.86,1.62,PINK,15.5)

    s = new_slide(prs,"原论文各用了什么数据？① 肝、鼠脑与乳腺","不能直接比较各论文表里的分数：组织、切片划分、预测基因数和评价方式都不相同。","原文：BLEEP §3.1；ResSAT 2026 Methods / additional evaluations；ST-Net 2020 摘要与 Methods。",
                  "BLEEP取四切片各自top1000高变基因并集，共3467，留出第3切片；本项目只在训练集选200基因。ResSAT正式版新增SCC和HER2ST，不能仍说只有两套鼠脑。ST-Net原数据与HER2ST不是同一个队列。SCC来源GSE144240；HER2ST来源github.com/almaan/her2st。")
    card(s,"BLEEP · 人肝","GSE240429：4 张连续切片\n9,269 个位置；Visium，55 μm\n原论文预测 3,467 个基因\n\n本地：四张 H&E、表达均已下载\n本次统一比较使用这套数据",.64,1.94,3.92,4.72,BLUE,16)
    card(s,"ResSAT · 鼠脑 + 人体组织","鼠脑前部 SA、后部 SP\n各 2 张；评测 2,000 个高变基因\n2026 正式版还选用：\nSCC（皮肤鳞癌）：12 张中用 3 张\nHER2ST（乳腺癌）：36 张中用 6 张\n\n本地：SA/SP 四张原图、计数、\n坐标已下载；后两套未完整下载",4.73,1.94,3.96,4.72,LIGHT,15.5)
    card(s,"ST-Net · 乳腺癌","23 位患者，68 张切片\n30,612 个位置；早期 ST 平台\n测量区域约 100 μm\n原文按患者留出测试\n\n本地：未下载完整原数据\n本次只做 ST-Net 风格 MLP\n不是原版 ST-Net 复现",8.87,1.94,3.82,4.72,PINK,16)

    s = new_slide(prs,"原论文各用了什么数据？② Stem 与 GenAR","两者覆盖多组织；通过 HEST 整理或获取配对的 H&E、空间表达矩阵及坐标。","原文：Stem §5.1–5.2 / 附录 D；GenAR §4.1–4.2 / 附录 A（本地 arXiv 版本）。",
                  "Kidney:Stem23切片22人，200基因，主测试20-0038(AKI)；GenAR测试NCBI697。HER2ST:Stem36切片8人，每片176–712spots，300 HMHVG或296 DEGs，测试B1；GenAR的HEST版共13594spots，测试SPA148。PRAD23切片2人，1418–4079spots/片，测试MEND145。健康鼠脑14切片4只雄鼠，2675–3617spots/片，测试NCBI667。GenAR额外ccRCC24片，测试INT2。不要未经核对把两版本样本ID强行对应。")
    columns=[(.68,2.54),(3.28,3.8),(7.15,5.46)]
    for (x,w),value in zip(columns,["数据集 / 组织","数据规模与内容","两篇论文怎么使用"]):
        rect(s,x,1.86,w,.47,RED); txt(s,value,x+.1,1.9,w-.18,.36,17,"FFFFFF",True)
    rows=[
        ("Kidney Visium\n人肾脏","23 张；健康 / 慢性肾病 / 急性肾损伤\n315–4,159 spots/张；55 μm","Stem 主实验：200 基因\nGenAR 主实验：200 基因，测试 NCBI697"),
        ("HER2ST\nHER2 阳性乳腺癌","原数据 36 张 / 8 人；100 μm\nGenAR 所用版本共 13,594 spots","Stem：300 高表达高变 / 296 差异基因\nGenAR：200 基因，测试 SPA148"),
        ("PRAD\n人前列腺癌","23 张 / 2 人；Visium，55 μm\n1,418–4,079 spots/张","Stem 附加实验；GenAR 主实验\n均报告测试切片 MEND145"),
        ("Healthy Mouse Brain\n健康小鼠脑","14 张 / 4 只小鼠；Visium，55 μm\n2,675–3,617 spots/张","Stem 附加实验；GenAR 主实验\n均报告测试切片 NCBI667"),
    ]
    for i,row in enumerate(rows):
        y=2.43+i*.82
        for (x,w),value in zip(columns,row):
            rect(s,x,y,w,.76,LIGHT if i%2==0 else BLUE); txt(s,value,x+.1,y+.085,w-.2,.64,14.7)
    txt(s,"补充：GenAR 还在附录验证了 ccRCC（肾透明细胞癌，24 张切片）。",.74,5.87,11.9,.4,16)
    rect(s,.67,6.36,12.0,.46,PINK)
    txt(s,"本地未下载上述完整队列；已下载的 HEST-IDC 四切片是另一项乳腺外部核验，不能混称为这些数据。",.79,6.41,11.77,.33,14.8,RED,True)

    s = new_slide(prs,"同一测试集上，各方法究竟表现怎样？","A1+B1 训练 → C1 验证 → D1 测试；固定 200 基因。以下全部从已有预测重新计算，未重新训练。","数据与逐基因指标：results/benchmark_ppt_audit.json；旧版混合口径表不再用于这张比较。",
                  "统一评价：GenAR用原始预测计数；其他方法用max(expm1(log预测),0)；真值用原始200列计数。各自log1p(10000*c/sum200 c)，没有借用测试标签总计数。此为200基因面板相对表达，不是全转录组计数。PCC逐基因跨位置计算，对预先固定200基因平均，不能测试后挑top10。Stem为3次生成均值。Ridge/MLP/BLEEP/GenAR/Stem用冻结ResNet18；ResSAT微调ResNet50，预算也不同。")
    txt(s,"平均 PCC ↑（越接近 1 越好）",.68,1.82,6.7,.4,20,RED,True)
    txt(s,"MAE ↓：平均差多少",7.46,1.82,2.79,.4,20,RED,True)
    txt(s,"数值不代表原论文排名",10.27,1.85,2.5,.4,14,GRAY)
    order=sorted(stats,key=lambda n:stats[n]["pcc"],reverse=True)
    comments={"ResSAT":"本次最高，仍是弱相关","BLEEP":"略高于 MLP / Ridge","MLP 基线":"简单图像回归","Image Ridge":"低复杂度参照","GenAR":"平均相关接近 0","Stem":"未恢复逐基因空间信号"}
    for i,name in enumerate(order):
        y=2.44+i*.46
        txt(s,name,.74,y,1.68,.37,18,RED if name=="ResSAT" else INK,name=="ResSAT")
        rect(s,2.52,y+.06,stats[name]["pcc"]/.10*3.69,.22,RED if name=="ResSAT" else "819CAD")
        txt(s,f"{stats[name]['pcc']:.4f}",6.13,y-.01,1.08,.37,18,bold=True)
        txt(s,f"{stats[name]['mae']:.3f}",7.5,y-.01,1.1,.37,18)
        txt(s,comments[name],9.08,y+.005,3.44,.36,15)
    rect(s,.67,5.43,12.0,.53,LIGHT)
    txt(s,f"只输出训练均值、不看图像：MAE = {audit['mean_baseline']['mae']:.3f}，PCC 无定义。ResSAT 的 MAE 仅略低于它。",.83,5.53,11.67,.34,16,RED,True)
    txt(s,"PCC：逐基因比较空间变化，再平均 200 个基因；MAE：所有预测值的平均绝对误差。\n统一单位：200 基因合计缩放至 10,000，再取 log1p；预测与真值各自换算。\n比较边界：本项目适配版本；图像编码器、训练预算不同，不是原论文方法的公平排名。",.72,6.07,11.96,.96,15)

    examples=[
        ("GLUL","预测可视化① GLUL：是否找到了相似的高表达区域？","Ridge / MLP / BLEEP / ResSAT 比实测更平滑：抓住部分区域趋势，但丢失了许多局部差异。\nStem 颜色变化丰富，却未出现在正确位置（PCC 0.018）；“看起来丰富”不等于预测准确。"),
        ("IGKC","预测可视化② IGKC：换一个基因，差异还成立吗？","GenAR 的大面积低值区与实测分布不对应（PCC −0.090）；不能仅凭整张图的颜色深浅判断效果。\nBLEEP、MLP、ResSAT 在该基因上均约 0.17，无明显领先者；总排名仍应看第 6 页的全基因平均。"),
    ]
    for gene,heading,analysis in examples:
        s=new_slide(prs,heading,"同一 D1 切片、同一批 2,265 个位置；每个点是一个测量区域。先看“实测真值”，再横向比较模型。","真实预测 NPZ 重绘；同页共享色标 / 坐标范围，无插值与平滑。PCC 以全部位置的连续数值计算。",
                    "GLUL与IGKC沿用此前已展示的示例，非据新表挑最佳基因。每页色标上限取真值及六法预测的共同最大值，下限0，不截断。不同基因页允许不同上限。图中PCC与本轮审计JSON一致。\n"+analysis)
        picture(s,ASSETS/f"predictions_{gene}.png",.54,1.82,12.27,4.37)
        rect(s,.64,6.20,12.05,.78,LIGHT,True); txt(s,analysis,.81,6.27,11.72,.66,15)

    s=new_slide(prs,"本轮更新了什么？目前能得出什么结论？","本次汇报聚焦一套主测试数据：把“数据在哪里、比较什么、结果怎么看”放在同一条线上。","所有重新评测均复用现有权重与预测；未新增训练。原版 PPT 保留；本次新增 9 页统一 Benchmark 汇报。",
                "这是阶段性证据，不宣称已经完全复现成功。优先补齐同编码器/预算对照、生成式模型适配诊断，再扩展独立患者。同一供者连续切片测试不能证明跨患者泛化。原论文复现和HEST-IDC外部核验不与此次主测试混排。")
    card(s,"已完成 · 可以直接检查","四张 H&E、表达、坐标与划分逐项核对。\n纠正旧表混用归一化口径的问题，六法重算。\n两组基因展示六种方法，统一色标重绘。",.67,1.95,12.0,2.00,BLUE,17)
    card(s,"当前结论 · 不夸大效果","ResSAT 在本次适配中平均 PCC 最高，但仅为 0.087，预测能力仍有限。\nRidge / MLP / BLEEP 提供可用参照；GenAR / Stem 尚未显示整体优势。\n这不说明原论文方法无效，也不能证明对其他患者、器官同样成立。",.67,4.02,12.0,1.93,LIGHT,17)
    txt(s,"下一步：统一编码器与训练预算 → 排查生成式适配 → 增加独立患者测试",.76,6.16,11.93,.37,19,RED,True)
    txt(s,"请老师批评指正",.78,6.62,11.88,.44,21,RED,True,PP_ALIGN.CENTER)
    assert len(prs.slides)==9
    prs.core_properties.title="H&E 空间转录组生成研究：统一 Benchmark 汇报"
    prs.core_properties.author="庹力元"
    prs.core_properties.subject="数据来源、统一面板评价和逐基因空间预测对比"
    prs.save(OUT)
    print(f"PPT: {OUT}\n页数: {len(prs.slides)}")


if __name__=="__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    truth,genes,xy,predictions,audit=load_and_audit()
    make_figures(truth,genes,xy,predictions,audit)
    make_deck(truth,genes,xy,predictions,audit)
    for name,result in audit["stats"].items():
        print(name,f"PCC={result['pcc']:.6f}",f"MAE={result['mae']:.6f}")
