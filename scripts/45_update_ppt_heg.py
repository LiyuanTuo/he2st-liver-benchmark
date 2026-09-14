"""在原统一 Benchmark PPT 中插入 2 页（原因分析 + HEG 改进结果），共 11 页。

不改动旧 9 页的内容结构，仅：
1. 在第 6 页（旧口径结果表）后插入第 7、8 页；
2. 更新第 6 页提示、第 9/10 页（示例）的页码引用、第 11 页（结论）的表述；
3. 全部页面重新编号为 01-11 / 11。

数值来自 results/heg_benchmark_final.json（由 scripts/44 生成）。
运行：py -3.12 scripts/45_update_ppt_heg.py
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "H&E空间转录组生成研究_统一Benchmark汇报_庹力元.pptx"
ASSETS = ROOT / "figures/benchmark_ppt"
DATA = json.loads((ROOT / "results/heg_benchmark_final.json").read_text(encoding="utf-8"))

RED, GOLD, INK, GRAY = "B71C1C", "B08D57", "262626", "666666"
LIGHT, BLUE, GREEN, PINK = "F7F4EF", "EDF3F8", "EDF5EF", "F9EDEC"

HEG_RAW = DATA["heg_panel"]["raw"]
HEG_RECON = DATA["heg_panel"]["pca50_reconstructed"]
OLD_RAW = DATA["old_panel_raw_all200"]
ORDER = sorted(HEG_RAW, key=lambda n: (HEG_RECON[n]["top50_heg_train_selected"] or -1), reverse=True)


def rect(slide, x, y, w, h, fill=LIGHT, rounded=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = .06
    return shape


def txt(slide, value, x, y, w, h, size=18, ink=INK, bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(.015)
    frame.margin_top = frame.margin_bottom = Inches(.02)
    for i, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line; p.font.name = "Microsoft YaHei"; p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(ink); p.font.bold = bold
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(4); p.line_spacing = 1.12
    return shape


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - iw * scale) / 2), Inches(y + (h - ih * scale) / 2),
        width=Inches(iw * scale), height=Inches(ih * scale),
    )


def insert_slide(prs, index):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[-1])
    xml_slides.insert(index, slides[-1])
    return slide


def new_slide_style(prs, index, heading, subtitle, source, notes):
    s = insert_slide(prs, index)
    s.background.fill.solid(); s.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    txt(s, heading, .48, .26, 12.25, .51, 27, RED, True)
    rect(s, .5, .94, 12.3, .025, GOLD)
    txt(s, subtitle, .51, 1.07, 12.24, .57, 16, GRAY)
    txt(s, source, .5, 7.12, 11.86, .22, 9, GRAY)
    txt(s, "00 / 11", 12.0, 7.11, .79, .24, 10, GRAY, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    return s


def replace_in_slide(slide, old, new):
    replaced = False
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if old in run.text:
                    run.text = run.text.replace(old, new)
                    replaced = True
    return replaced


def replace_anywhere(prs, old, new, limit=None):
    count = 0
    for slide in prs.slides:
        if replace_in_slide(slide, old, new):
            count += 1
            if limit and count >= limit:
                break
    return count


def build_slide_cause(prs):
    values = DATA["heg_panel"]
    ressat_old = OLD_RAW["ResSAT"]
    ressat_heg_raw = values["raw"]["ResSAT"]["all_200"]
    ressat_heg_top50 = values["pca50_reconstructed"]["ResSAT"]["top50_heg_train_selected"]
    s = new_slide_style(
        prs, 6,
        "统一基准的 PCC 为什么只有 0.0x？",
        "同一个 ResSAT 模型实现已在官方数据上复现论文数值（top-50 HEG：0.880 vs 论文 0.878）；"
        "0.0x 来自三个协议差异叠加，而不是模型损坏。",
        "证据：results/ressat_official/evaluation.json；results/heg_benchmark_final.json；"
        "data/processed/gse240429(_heg)/manifest.json",
        "三个原因各自有量化证据：① 旧 200 基因面板按方差选择，与全转录组 top-50 高表达基因零重叠；"
        "② 论文在 PCA/Harmony 重建空间评测，旧基准在完整 log 空间评测含噪声真值；"
        "③ 旧基准是跨切片划分且各方法训练预算不同。改进后表格见第 8 页。",
    )
    cards = [
        ("① 基因面板选错了", BLUE,
         f"旧面板按方差选 200 基因，全转录组 top-50 高表达基因（ALB、HP、SAA1…）一个都不在面板内"
         f"（重叠 0/50）；面板内低表达基因的逐基因 PCC 被噪声主导，把平均拉到 {ressat_old:.3f}。\n"
         "高表达基因表达量高、技术噪声占比小，是论文 headline 用的口径。"),
        ("② 评价空间不同", GREEN,
         f"论文对真值与预测同时做 PCA/Harmony 重建（去噪）后算 PCC；旧基准在完整 log 空间评原始噪声。\n"
         f"同一批旧预测仅换评价空间：{ressat_old:.3f} → 0.154（全 200 基因），噪声成分被去除后显著上升。"),
        ("③ 划分与预算更硬", PINK,
         "论文：鼠脑相邻切片、2,000 HVG、病理预训练 ResNet50 微调；\n"
         "旧基准：人肝 A/B → D 跨切片（域偏移大）、200 基因、冻结 ResNet18 或轻量微调。\n"
         "本地官方案例 top-50 HEG 已复现 0.880，说明实现无系统性错误。"),
    ]
    for i, (heading, fill, body) in enumerate(cards):
        rect(s, .66 + i * 4.12, 1.86, 3.96, 2.62, fill, True)
        txt(s, heading, .85 + i * 4.12, 2.0, 3.6, .44, 19, RED, True)
        txt(s, body, .85 + i * 4.12, 2.56, 3.6, 1.84, 13.5)
    picture(s, ASSETS / "cause_decomposition_v2.png", .62, 4.62, 12.14, 2.34)


def build_slide_improvement(prs):
    s = new_slide_style(
        prs, 7,
        "改进：重选高表达基因面板，按论文口径重训与重评",
        "top-200 HEG 面板（按训练切片平均表达选，含 ALB、HP、RPS/RPL 等）；"
        "六种方法同面板重训；top-50 HEG 在 PCA-50 重建空间评测（与论文同口径）。",
        "证据：results/heg_benchmark_final.json；logs_heg_pipeline.txt；"
        "data/processed/gse240429_heg/manifest.json（基因列表 SHA256 已记录）",
        "面板重选只改基因选择标准（高表达代替高方差），划分、图像块、编码器特征、barcode 行序不变。"
        "训练集选 HEG 无测试泄漏。剩余差距主要来自跨切片域偏移与完整空间评价。",
    )
    # 表头
    headers = [("方法", 1.05), ("旧面板 · 原空间\n全 200 基因", 3.15), ("HEG 面板 · 原空间\n全 200 基因", 3.15),
               ("HEG 面板 · PCA-50 重建\ntop-50 HEG", 3.15)]
    x0 = .62
    for title, width in headers:
        rect(s, x0, 1.82, width, .62, RED)
        txt(s, title, x0 + .08, 1.87, width - .16, .52, 14.5, "FFFFFF", True, PP_ALIGN.CENTER)
        x0 += width
    y = 2.5
    for i, name in enumerate(ORDER):
        fill = LIGHT if i % 2 == 0 else BLUE
        x0 = .62
        values = [name, OLD_RAW.get(name), HEG_RAW[name]["all_200"],
                  HEG_RECON[name]["top50_heg_train_selected"]]
        formats = [f"{v}", f"{v:.3f}", f"{v:.3f}", f"{v:.3f}"]
        bold_last = name == ORDER[0]
        for j, (title, width) in enumerate(headers):
            rect(s, x0, y, width, .44, fill)
            value = values[j]
            text = formats[j].format(v=value) if value is not None else "—"
            txt(s, text, x0 + .08, y + .045, width - .16, .36, 14.5,
                RED if j == 3 else INK, bold=(j == 0) or (j == 3 and bold_last),
                align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
            x0 += width
        y += .47
    top = ORDER[0]
    txt(s, f"ResSAT：旧口径 {OLD_RAW['ResSAT']:.3f} → HEG 面板重建空间 top-50 "
           f"{HEG_RECON['ResSAT']['top50_heg_train_selected']:.3f}（其余方法同样提升）。",
        .68, y + .06, 12.0, .38, 16.5, RED, True)
    txt(s, "改进措施：① 面板改按高表达选择（top-200 HEG，零值率 1.4% vs 旧面板 24.7%）"
           " ② 全方法同面板重训（GenAR/BLEEP/Stem 用验证集选定的改进版）"
           " ③ 按论文口径评价（PCA-50 重建空间 + top-50 HEG，训练集选基因、无测试泄漏）。\n"
           "剩余差距来源：跨切片域偏移（A/B→D）与完整空间评价；"
           "本地官方案例同口径已达 0.880，说明实现无误，后续可做领域自适应与病理编码器。",
        .68, y + .52, 12.02, 1.06, 14.5)


def update_slide_six(prs):
    """第 6 页：把右上角灰字换成指向第 8 页的提示。"""
    slide = prs.slides[5]
    for shape in slide.shapes:
        if shape.has_text_frame and "数值不代表原论文排名" in shape.text_frame.text:
            shape.text_frame.clear()
            p = shape.text_frame.paragraphs[0]
            p.text = "旧面板口径 · HEG 口径见第 8 页"
            p.font.name = "Microsoft YaHei"; p.font.size = Pt(14)
            p.font.color.rgb = RGBColor.from_string(GRAY); p.alignment = PP_ALIGN.RIGHT
            return


def update_conclusion(prs):
    slide = prs.slides[-1]
    ressat = HEG_RECON["ResSAT"]["top50_heg_train_selected"]
    old = OLD_RAW["ResSAT"]
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        text = shape.text_frame.text
        if text.startswith("ResSAT 在本次适配中平均 PCC 最高"):
            frame = shape.text_frame
            frame.clear()
            paragraphs = [
                f"旧面板原空间：ResSAT 平均 PCC {old:.3f}（六法最高），噪声主导、明显偏低。",
                f"改按论文口径（高表达基因 + PCA 重建空间）重选面板并重训后，top-50 HEG 平均 PCC 达 {ressat:.3f}。",
                "剩余差距主要来自跨切片划分与完整空间评价；本地官方案例已复现 0.880，说明模型实现无误。",
            ]
            for i, line in enumerate(paragraphs):
                p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
                p.text = line
                p.font.name = "Microsoft YaHei"; p.font.size = Pt(17)
                p.font.color.rgb = RGBColor.from_string(INK)
                p.space_before = Pt(0); p.space_after = Pt(4); p.line_spacing = 1.12
            return


def renumber(prs):
    total = len(prs.slides)
    pattern = re.compile(r"^\d{2} / \d{2}$")
    for index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if shape.has_text_frame and pattern.match(shape.text_frame.text.strip()):
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.text = f"{index:02d} / {total:02d}"
                p.font.name = "Microsoft YaHei"; p.font.size = Pt(10)
                p.font.color.rgb = RGBColor.from_string(GRAY)
                p.alignment = PP_ALIGN.RIGHT


def main():
    prs = Presentation(str(PPTX))
    assert len(prs.slides) == 9, f"期望 9 页旧 PPT，实际 {len(prs.slides)} 页"
    build_slide_cause(prs)
    build_slide_improvement(prs)
    update_slide_six(prs)
    replace_anywhere(prs, "总排名仍应看第 6 页的全基因平均。",
                     "总排名仍应看第 6 页（旧面板）与第 8 页（HEG 面板）的全基因平均。")
    replace_anywhere(prs, "所有重新评测均复用现有权重与预测；未新增训练。原版 PPT 保留；本次新增 9 页统一 Benchmark 汇报",
                     "新增 HEG 面板重训与论文口径重评；旧面板基准与官方案例结果全部保留。原版 PPT 保留；本次新增 11 页统一 Benchmark 汇报")
    update_conclusion(prs)
    renumber(prs)
    assert len(prs.slides) == 11, f"期望 11 页，实际 {len(prs.slides)} 页"
    prs.core_properties.title = "H&E 空间转录组生成研究：统一 Benchmark 汇报（HEG 改进版）"
    prs.save(PPTX)
    print(f"saved {PPTX}, slides={len(prs.slides)}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
