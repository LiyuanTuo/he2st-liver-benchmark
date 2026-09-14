"""Update the user's original PPT in place; retain 11 pages and prior reproduction."""
from pathlib import Path
import importlib.util,json,shutil
from pptx import Presentation
from pptx.enum.text import PP_ALIGN
ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('verified',ROOT/'scripts/53_update_verified_ppt.py')
style=importlib.util.module_from_spec(spec);spec.loader.exec_module(style)
OUT=ROOT/'results/liver_context_20260913';FIG=ROOT/'figures/liver_context_20260913'
txt,rect,card,header,pic=style.txt,style.rect,style.card,style.header,style.pic
def read(p):return json.loads((ROOT/p).read_text(encoding='utf8'))
def main():
    result=read('results/liver_context_20260913/comparison.json');metrics=result['metrics'];selected=result['selected_by_C1'];best=metrics[selected]
    backup=OUT/'ppt_before_context';backup.mkdir(exist_ok=True)
    for suffix in ['.pptx','.pdf']:
        source=style.PPT.with_suffix(suffix)
        if source.exists() and not (backup/source.name).exists():shutil.copy2(source,backup/source.name)
    prs=Presentation(style.PPT);assert len(prs.slides)==11
    names=['Image Ridge','MLP','BLEEP','ResSAT','GenAR','Stem','ST-Net (DenseNet121)','ContextFusion-imagenet','ContextFusion-scratch','ContextFusion-multiseed']
    labels={n:n for n in names};labels.update({'ST-Net (DenseNet121)':'ST-Net / DenseNet121','ContextFusion-imagenet':'自建三尺度 / 预训练','ContextFusion-scratch':'自建三尺度 / 全随机','ContextFusion-multiseed':'自建三尺度 / 多种子选择'})
    s=prs.slides[5]
    header(s,6,'统一人肝 Benchmark：七种基线与自建模型',
      '同一 D1、训练选定 HEG200 / HEG50、同一未重建真值；新配置只按 C1 选择。',
      '来源：results/liver_context_20260913/combined_benchmark.csv；单视野与冻结特征消融见第9页及完整报告。',
      json.dumps(result,ensure_ascii=False)+'\n多种子行在seed42/17/83及等权均值中由C1选择，不能称所有重复都达到该值。新模型直接回归原面板log单位，不改变测试真值。不同方法的计算预算与输入视野仍有差异。')
    cols=[.67,4.77,6.94,9.05,10.62]
    for x,w,value in zip(cols,[4.0,2,2,1.4,1.55],['方法','HEG200 PCC','HEG50 PCC','MAE','>0.5 基因']):txt(s,value,x,1.87,w,.4,14,style.RED,True)
    for i,name in enumerate(names):
        m=metrics[name];y=2.36+i*.355;rect(s,.6,y-.025,12.08,.35,'EDF5EF' if i>=7 else ('F7F4EF' if i%2==0 else 'FFFFFF'))
        for x,w,value in zip(cols,[4.0,2,2,1.4,1.55],[labels[name],f'{m["HEG200"]:.4f}',f'{m["HEG50"]:.4f}',f'{m["MAE"]:.3f}',f'{m["genes_above_05"]}/200']):txt(s,value,x,y,w,.31,13,style.INK,name==selected)
    txt(s,'“单基因超过0.5”与“固定面板平均超过0.5”是两件事；完整逐基因结果全部保留。',.67,6.18,12,.52,17,style.RED,True)
    s=prs.slides[6]
    header(s,7,'低 PCC 的原因：已纠正的问题与本轮证据',
      '历史错误已经修复；剩余差距需用同条件实验解释，不能继续归咎于低表达基因。',
      '证据：docs/05_PCC_DIAGNOSIS.md、docs/06_LIVER_CONTEXT_IMPROVEMENT.md；所有控制与失败尝试均保存。',
      '旧面板错误与Stem/GenAR/计分问题见05文档。当前面板固定不再更换。邻居RNA诊断使用验证集实测表达，不能作H&E预测成绩；半计数拆分不是技术重复，也不是性能硬上限。\n同数据研究：https://link.springer.com/article/10.1186/s12859-026-06447-7；HEG 0.310，不同split/Harmony/gene panel，不可直接比数值。')
    card(s,.6,1.93,5.92,2.15,'① 高表达面板：已修正，不再换',
      '旧面板与 HEG200 重叠为0；现为训练集选基因。\nD1 零值率已从24.70%降至3.45%。\n当前低分不能全归因于之前选错基因。','F9EDEC',15)
    card(s,6.72,1.93,5.98,2.15,'② 训练目标与最终计分需要对齐',
      '旧模型学全转录组分母，评测却换200基因分母。\n同多尺度 Ridge：D1 PCC 0.1018→0.1192。\n修正确有帮助，但不足以单独解决差距。','EDF3F8',15)
    card(s,.6,4.29,5.92,2.27,'③ 能拟合训练信号，泛化仍不足',
      '三视野预训练模型的 HEG50 PCC：\n训练0.608 → C1验证0.485 → D1测试0.302。\n只加冻结特征回归头，未超过原ST-Net。\n额外视野与新训练配方的作用见第9页。','EDF5EF',15)
    card(s,6.72,4.29,5.98,2.27,'④ 计数差异有贡献，但并非全部',
      '每spot UMI中位数：C1 12423；D1 6010。\nC1降采样诊断：HEG200 0.348→0.303。\n高表达仍不等于强空间变化、充足计数。\n这些因素未能解释全部差距，非性能硬上限。','F7F4EF',15)
    # Keep page 8: original ResSAT numerical reproduction and metric caveats.
    s=prs.slides[8]
    header(s,9,'自建 ContextFusion：设计、从零训练与消融',
      '58 / 231 / 461 μm → 共享 ResNet18 → 拼接 + 256维非线性头 → 200基因面板 log 表达。',
      '代码：scripts/59、63、64、65；C1选epoch/TTA/种子聚合，D1只作最终评测。全随机组不加载任何预训练权重。',
      '所有CNN参数可训练。训练标准化统计仅来自A1+B1，MSE+0.2*(1-PCC)，最多60轮，patience12。单视野控制把224px图像重复三次，保持参数量/损失/优化器等一致。比较不是ResNet18原论文复现，而是本项目新设计的适配模型。')
    pic(s,FIG/'learning_curves.png',.62,1.89,7.0,3.20)
    txt(s,'预训练与全随机版本均已完成训练，曲线仅为C1。',.78,5.16,6.7,.45,14,style.GRAY)
    entries=[('SingleView-control','同配方单视野'),('ContextFusion-imagenet','三视野 / 预训练'),('ContextFusion-scratch','三视野 / 全随机'),('ContextFusion-multiseed','三视野 / 多种子')]
    txt(s,'D1 消融：HEG200 / HEG50',7.85,1.92,4.7,.48,17,style.RED,True)
    for i,(key,label) in enumerate(entries):
        m=metrics[key];y=2.61+i*.73
        txt(s,label,7.9,y,4.5,.35,15,style.INK,True)
        txt(s,f'{m["HEG200"]:.4f} / {m["HEG50"]:.4f}',7.9,y+.36,4.5,.31,15,style.RED)
    card(s,.62,5.76,12.04,.99,'解释提升需要对照',
      '同配方单视野 0.1941/0.2790 → 三视野 0.2176/0.3019；其余提升不能全部归于视野。','EDF3F8',13)
    s=prs.slides[9]
    header(s,10,'固定实例：ALB 与 HP，继续按训练表达排名选例',
      '训练表达第1、2位；展示真值、旧 ST-Net、C1选定新模型及全随机版本，同一坐标与色标。',
      '来源：results/liver_context_20260913/combined_predictions.npz；色标裁剪只用于显示，不改变PCC。',
      '没有按测试PCC选择实例。单基因预测好坏与面板平均分别报告。图示是表达预测，不能替代RNA测序。')
    pic(s,FIG/'ALB_HP_fixed_examples.png',.58,1.85,12.14,4.84)
    s=prs.slides[10]
    threshold='尚未证明固定面板平均达到0.5。' if best['HEG200']<.5 and best['HEG50']<.5 else '具体超过0.5的面板见本页数值。'
    base=metrics['ST-Net (DenseNet121)']
    header(s,11,'当前结论：可学习信号更明确，仍需正视剩余差距',
      f'C1选定：{selected}；{threshold}',
      '完整报告：docs/06_LIVER_CONTEXT_IMPROVEMENT.md；预测、权重、消融、失败控制和验证记录均可复核。',
      json.dumps(result,ensure_ascii=False)+'\n空间bootstrap区间只针对一张切片，不是跨供者泛化区间。0.5目标未达到时如实报告，不能换测试基因或平滑真值来制造达标。')
    card(s,.62,1.95,12.05,1.55,'同一测试任务的实际变化',
      f'HEG200：旧 ST-Net {base["HEG200"]:.4f} → 新模型 {best["HEG200"]:.4f}；HEG50：{base["HEG50"]:.4f} → {best["HEG50"]:.4f}。\n新模型有 {best["genes_above_05"]}/200 个基因 PCC>0.5，其中固定 HEG50 内有 {best["HEG50_above_05"]}/50 个。','EDF5EF',17)
    card(s,.62,3.75,5.92,2.15,'现在能支持的结论',
      '部分基因存在可学习的H&E形态信号。\n新模型与全随机组均已实际训练、测试。\n所有固定基因完整计分，没有只报高分基因。\n模型差距用同条件消融而非猜测解释。','F7F4EF',15)
    card(s,6.73,3.75,5.92,2.15,'现在不能支持的结论',
      '不能称全部高表达基因都预测得好。\n不能将鼠脑重建目标的0.698要求于此任务。\n单供者连续切片不足以验证跨患者泛化。\n需要更多独立供者及匹配的病理编码器。','EDF3F8',15)
    txt(s,'用同一目标上的真实提升证明进步；达不到0.5也保留完整证据。',.76,6.26,11.8,.5,20,style.RED,True)
    for slide in list(prs.slides)[:5]:
        notice='本轮续研：七种原基线之外新增自建ContextFusion及控制；结果见第6、9、11页，基因/切片/真值保持不变。'
        if notice not in style.note_frame(slide).text:style.note_frame(slide).text+='\n'+notice
    style.style.renumber(prs);prs.core_properties.subject='2026-09-13 liver context model follow-up; 11 slides';prs.save(style.PPT)
    check=Presentation(style.PPT);assert len(check.slides)==11
    problems=[]
    for i,slide in enumerate(check.slides,1):
        for shape in slide.shapes:
            if i>=6 and (shape.left<0 or shape.top<0 or shape.left+shape.width>check.slide_width+1000 or shape.top+shape.height>check.slide_height+1000):problems.append([i,shape.name])
    assert not problems,problems
    (OUT/'ppt_verification.json').write_text(json.dumps(dict(slides=11,edited=[6,7,9,10,11],original_file=str(style.PPT.name),backup=str(backup.relative_to(ROOT))),ensure_ascii=False,indent=2),encoding='utf8')
    print(style.PPT)
if __name__=='__main__':main()
