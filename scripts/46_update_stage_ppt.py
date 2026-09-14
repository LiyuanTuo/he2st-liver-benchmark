"""把 HEG 改进结果加入《阶段汇报》PPT（保持原版风格）。

在原 PPT 倒数第二页（结束页之前）插入两页：
- 22 原因量化：统一 benchmark 的 PCC 为什么只有 0.0x
- 23 改进结果：重选高表达基因（HEG）面板并重训全部方法

风格与原版一致：微软雅黑、B71C1C 深红标题、B08D57 金色分隔线、
浅色卡片 + 左侧色条（3F5F89 / 40765A / B71C1C）、圆角 0.08、描边 D8D2C8。
数值来自 results/heg_benchmark_final.json（由 scripts/44 生成）。

运行：py -3.12 scripts/46_update_stage_ppt.py
"""

from __future__ import annotations

import json
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "H&E空间转录组生成研究_阶段汇报_庹力元.pptx"
ASSETS = ROOT / "figures/benchmark_ppt"
DATA = json.loads((ROOT / "results/heg_benchmark_final.json").read_text(encoding="utf-8"))

RED = "B71C1C"; GOLD = "B08D57"; INK = "262626"; GRAY = "666666"
BLUE, GREEN, PINK = "3F5F89", "40765A", "B71C1C"
L_BLUE, L_GREEN, L_PINK, L_WHITE = "EAF2F8", "EAF4EE", "F8ECEB", "FFFFFF"
BORDER = "D8D2C8"
FONT = "Microsoft YaHei"

HEG_RAW = DATA["heg_panel"]["raw"]
HEG_RECON = DATA["heg_panel"]["pca50_reconstructed"]
OLD_RAW = DATA["old_panel_raw_all200"]
ORDER = sorted(HEG_RAW, key=lambda n: (HEG_RECON[n]["top50_heg_train_selected"] or -1), reverse=True)


def txt(slide, value, x, y, w, h, size=12, ink=INK, bold=False, align=PP_ALIGN.LEFT, spacing=1.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.word_wrap = True; frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.margin_left = frame.margin_right = Inches(.02)
    frame.margin_top = frame.margin_bottom = Inches(.01)
    for i, line in enumerate(str(value).split("\n")):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = line
        p.font.name = FONT; p.font.size = Pt(size)
        p.font.color.rgb = RGBColor.from_string(ink); p.font.bold = bold
        p.alignment = align
        p.space_before = Pt(0); p.space_after = Pt(2); p.line_spacing = spacing
    return shape


def rect(slide, x, y, w, h, fill, rounded=False, border=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h),
    )
    if rounded:
        shape.adjustments[0] = .08
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor.from_string(fill)
    if border:
        shape.line.color.rgb = RGBColor.from_string(border); shape.line.width = Pt(.8)
    else:
        shape.line.fill.background()
    return shape


def header(slide, section, title, page):
    txt(slide, section, .55, .13, 2.4, .22, 8.5, RED, True)
    txt(slide, title, .55, .38, 11.9, .48, 24, RED, True)
    rect(slide, .55, .93, 12.22, .03, GOLD)
    txt(slide, str(page), 12.35, .16, .42, .22, 9, GRAY, True, PP_ALIGN.RIGHT)


def card(slide, x, y, w, h, fill, accent, title, body, body_size=11.5):
    rect(slide, x, y, w, h, fill, rounded=True, border=BORDER)
    rect(slide, x, y, .06, h, accent)
    txt(slide, title, x + .2, y + .12, w - .4, .32, 13.5, accent, True)
    txt(slide, body, x + .2, y + .52, w - .4, h - .62, body_size, INK)


def picture(slide, path, x, y, w, h):
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    return slide.shapes.add_picture(
        str(path), Inches(x + (w - iw * scale) / 2), Inches(y + (h - ih * scale) / 2),
        width=Inches(iw * scale), height=Inches(ih * scale),
    )


def insert_slide(prs, index):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    xml_slides = prs.slides._sldIdLst
    slides = list(xml_slides)
    xml_slides.remove(slides[-1])
    xml_slides.insert(index, slides[-1])
    return slide


def build_cause_slide(prs):
    slide = insert_slide(prs, len(prs.slides) - 1)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    header(slide, "08 实验结果", "统一 benchmark 的 PCC 为什么只有 0.0x？——原因量化", 22)
    ressat_old = OLD_RAW["ResSAT"]
    cards = [
        ("① 基因面板选错：按方差选，不含高表达基因", L_PINK, PINK,
         f"旧 200 基因面板按方差选择，全转录组 top-50 高表达基因（ALB、HP、SAA1…）"
         f"一个都不在面板内（重叠 0/50）；低表达基因的逐基因 PCC 被噪声主导，"
         f"把宏平均拉到 {ressat_old:.3f}。"),
        ("② 评价空间不同：论文在 PCA 重建空间评测", L_BLUE, BLUE,
         "论文对真值与预测同时做 PCA-50 重建（去噪）后算 PCC；旧基准在完整 log 空间"
         "评原始噪声。同一批旧预测仅换评价空间：0.087 → 0.154（全 200 基因）。"),
        ("③ 数据与划分更难：跨切片 + 统一预算", L_GREEN, GREEN,
         "论文：鼠脑相邻切片、2,000 HVG、病理预训练 ResNet50；本项目：人肝 A/B→D 跨切片、"
         "200 基因、冻结 ResNet18。本地官方案例已复现 0.880/0.878，实现无系统性错误。"),
    ]
    for i, (title, fill, accent, body) in enumerate(cards):
        card(slide, .55 + i * 4.15, 1.12, 3.98, 2.6, fill, accent, title, body, 11)
    picture(slide, ASSETS / "cause_decomposition_v2.png", .55, 3.92, 12.22, 3.05)
    txt(slide, "结论：0.0x 来自三个协议差异叠加（面板 × 空间 × 划分），而不是模型损坏。",
        .55, 7.02, 12.2, .3, 12, RED, True)


def build_improvement_slide(prs):
    slide = insert_slide(prs, len(prs.slides) - 1)
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
    header(slide, "08 实验结果", "改进：重选高表达基因（HEG）面板，按论文口径重训与重评", 23)
    headers = [("方法", 2.6), ("旧面板 · 原空间\n全 200 基因", 3.2),
               ("HEG 面板 · 原空间\n全 200 基因", 3.2), ("HEG 面板 · PCA-50 重建\ntop-50 HEG", 3.22)]
    x0 = .55
    for title, width in headers:
        rect(slide, x0, 1.10, width, .5, RED)
        txt(slide, title, x0 + .06, 1.13, width - .12, .44, 12, "FFFFFF", True, PP_ALIGN.CENTER)
        x0 += width
    y = 1.66
    for i, name in enumerate(ORDER):
        fill = L_WHITE if i % 2 == 0 else L_BLUE
        x0 = .55
        values = [name, OLD_RAW.get(name), HEG_RAW[name]["all_200"],
                  HEG_RECON[name]["top50_heg_train_selected"]]
        for j, (title, width) in enumerate(headers):
            rect(slide, x0, y, width, .52, fill, border=BORDER)
            value = values[j]
            text = name if j == 0 else (f"{value:.3f}" if value is not None else "—")
            txt(slide, text, x0 + .1, y + .09, width - .2, .36, 12.5,
                RED if j == 3 else INK, bold=(j == 3), align=PP_ALIGN.CENTER if j else PP_ALIGN.LEFT)
            x0 += width
        y += .55
    ressat = HEG_RECON["ResSAT"]["top50_heg_train_selected"]
    card(slide, .55, 5.06, 6.0, 1.8, L_GREEN, GREEN, "改进措施",
         "① 面板改按高表达选择：top-200 HEG（含 ALB、HP 等），零值率 1.4% vs 旧面板 24.7%\n"
         "② 六种方法同面板重训，GenAR/Stem/BLEEP 用验证集选定版本\n"
         "③ 按论文口径评价：PCA-50 重建空间 + top-50 HEG（训练集选基因，无测试泄漏）", 10.5)
    card(slide, 6.77, 5.06, 6.0, 1.8, L_PINK, PINK, "改进后的结论",
         f"ResSAT 平均 PCC：{OLD_RAW['ResSAT']:.3f} → {ressat:.3f}（top-50 HEG，重建空间），"
         "其余方法同样提升。\n"
         "剩余差距来自跨切片域偏移与完整空间评价；官方案例同口径 0.880，"
         "下一步做领域自适应与病理编码器。", 10.5)


def main():
    prs = Presentation(str(PPTX))
    before = len(prs.slides)
    build_cause_slide(prs)
    build_improvement_slide(prs)
    assert len(prs.slides) == before + 2
    # 结束页保持在最后一页
    prs.core_properties.title = "H&E 空间转录组生成研究：阶段汇报（含 HEG 改进）"
    prs.save(PPTX)
    print(f"saved {PPTX}, slides={len(prs.slides)}")


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main()
