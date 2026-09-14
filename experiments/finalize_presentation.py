"""Finalize author details, ensemble explanation and final spatial examples in-place."""
from pathlib import Path
import importlib.util
import shutil
from pptx import Presentation
from pptx.util import Inches,Pt

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('s',ROOT/'scripts/53_update_verified_ppt.py')
s=importlib.util.module_from_spec(spec);spec.loader.exec_module(s)


def main():
    backup=ROOT/'results/finalization_20260915';backup.mkdir(exist_ok=True)
    if not (backup/s.PPT.name).exists():shutil.copy2(s.PPT,backup/s.PPT.name)
    deck=Presentation(s.PPT);assert len(deck.slides)==11
    cover=deck.slides[0]
    for shape in cover.shapes:
        if not shape.has_text_frame:continue
        if '统一 Benchmark 实验汇报' in shape.text:
            shape.text='庹力元  ·  学号59S2311  ·  2026.09.15\n东南大学生命科学与技术学院  ·  生物、计算机双学位学士'
            shape.height=Inches(.65);shape.width=Inches(12.2)
            for p in shape.text_frame.paragraphs:
                for r in p.runs:r.font.name='Microsoft YaHei';r.font.size=Pt(14)
        if shape.text=='同一张图，七种方法':
            for p in shape.text_frame.paragraphs:
                for r in p.runs:r.text='统一基线与自建集成'
    slide=deck.slides[9]
    s.header(slide,10,'最终空间预测：ALB 与 HP',
             '沿用训练表达排名前两位的示例；真值、ST-Net、原集成和最终集成共用坐标与色标。',
             '来源：report/figures/spatial_examples.png；展示色标截断不改变计分；个别图例不替代固定200/50基因的宏平均。',
             '由experiments/build_report_assets.py从冻结预测生成。最终为75%原三种子均值+25%新三模型均值；主结果固定D1 HEG200=0.2343219242,HEG50=0.3179922544。')
    s.pic(slide,ROOT/'report/figures/spatial_examples.png',.58,1.8,12.2,4.88)
    slide=deck.slides[10]
    s.header(slide,11,'最终集成：六个模型、固定权重、可复现结果',
             '加权平均发生在同一面板log表达空间；不平均PCC分数，不按测试基因分别调权重。',
             '完整报告：report/main.pdf（LaTeX源码main.tex）；组成与SHA256：benchmarks/liver/final_ensemble.json。',
             '作者：庹力元，59S2311，东南大学生命科学与技术学院，生物、计算机双学位学士。最终六模型权重重新推理最大差0。D1重复用于项目开发，未构成独立患者外部测试。')
    s.card(slide,.60,1.94,4.0,2.1,'① 原模型组 O：75%',
           'ResNet18：seed42 / 17 / 83\n三者等权平均，再乘0.75。\n每个原模型最终权重=1/4。\n三尺度视野均端到端训练。','EDF3F8',14)
    s.card(slide,4.72,1.94,4.0,2.1,'② 新模型组 N：25%',
           'ResNet18 / EfficientNet-B0\n另一个ResNet18用训练标签平滑。\n三者等权平均，再乘0.25。\n每个新模型最终权重=1/12。','EDF5EF',14)
    s.card(slide,8.84,1.94,3.88,2.1,'③ C1锁定混合权重',
           '最终预测 = 0.75 O + 0.25 N\n比较新组25% / 50% / 75%。\nC1 HEG200最高者为25%。\n六模型均使用选定的原权重。','F7F4EF',14)
    s.card(slide,.60,4.34,5.96,2.15,'④ 效果：有增益，但不是全面提升',
           'D1 HEG200/50：0.2343 / 0.3180。\n较原集成 +0.0042 / +0.0024；MAE略变差。\nHEG50增益区间跨零；固定均值仍未达0.5。\n全随机单模型HEG50为0.3278，更高。','F9EDEC',14)
    s.card(slide,6.76,4.34,5.96,2.15,'⑤ 复现：直接从六份权重推理',
           'run.py ensemble --slides C73_D1 --output …\n2265×200预测逐元素重现，最大差=0。\n正式报告、源码、结果表与最终PPT一起记录。\n后续重点：独立供体、未查看的外部测试。','EDF3F8',14)
    deck.save(s.PPT)
    print('Finalized original PPT:',len(deck.slides),'pages')


if __name__=='__main__':main()
